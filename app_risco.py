"""Avaliação de risco de avaliação negativa (v1.1).

Uso: dê dois cliques no programa e escolha a planilha com os pedidos entregues
(CSV ou Excel), ou arraste a planilha para cima do programa.

A planilha pode vir em dois formatos:
  1. Planilha simples (recomendado): uma linha por pedido, com datas, valores e
     quantidades, usando os nomes de coluna que a empresa já usa. O programa
     reconhece os nomes mais comuns; quando não reconhece, pergunta qual coluna
     corresponde a cada informação e lembra da resposta nas próximas vezes.
  2. Formato completo: as 24 colunas que o modelo usa (como em exemplo_pedidos.csv).

O resultado é salvo ao lado da planilha, em Excel, com o final "_com_risco.xlsx".
"""
import difflib
import json
import os
import re
import sys
import unicodedata

import joblib
import numpy as np
import pandas as pd

# Importações explícitas: o modelo salvo usa estas partes do scikit-learn.
# Sem elas, o PyInstaller não percebe que precisa incluí-las no .exe.
import sklearn.pipeline  # noqa: F401
import sklearn.compose  # noqa: F401
import sklearn.preprocessing  # noqa: F401
import sklearn.ensemble  # noqa: F401
import sklearn._loss.loss  # noqa: F401
import sklearn.utils.validation  # noqa: F401
import openpyxl  # noqa: F401  (necessário para ler e gravar Excel)
from openpyxl.styles import Alignment, Font, PatternFill

VERSAO = "1.1"
CORTE = 0.8        # risco mínimo para o pedido entrar em "Precisam de atenção"
CORTE_MEDIO = 0.6  # a partir daqui o risco é considerado médio

CATEGORICAS = ["categoria", "estado_vendedor", "estado_cliente", "forma_pagamento"]
NUMERICAS = ["dias_entrega", "dias_prometidos", "dias_atraso", "atrasado", "horas_aprovacao",
             "dias_ate_transportadora", "mes_compra", "dia_semana_compra",
             "n_itens", "n_produtos", "n_vendedores", "valor_total", "frete_total", "frete_sobre_valor",
             "peso_total_g", "fotos_media", "tam_descricao_media", "n_parcelas", "n_formas_pagamento",
             "mesmo_estado"]
COLUNAS_MODELO = CATEGORICAS + NUMERICAS

# ---------------------------------------------------------------------------
# Informações da planilha simples: (chave, descrição, obrigatória, nomes comuns)
# ---------------------------------------------------------------------------
CAMPOS = [
    ("pedido", "Código do pedido", True,
     ["pedido", "order_id", "id_pedido", "codigo", "codigo_pedido", "cod_pedido", "numero_pedido",
      "num_pedido", "n_pedido", "no_pedido", "nro_pedido", "id", "pedido_id", "cod", "codigo_venda"]),
    ("data_compra", "Data da compra", True,
     ["data_compra", "data_pedido", "dt_compra", "dt_pedido", "comprado_em", "data_venda", "dt_venda",
      "order_purchase_timestamp", "data", "data_emissao"]),
    ("data_prometida", "Data prometida de entrega", True,
     ["data_prometida", "prazo", "prazo_entrega", "data_prevista", "previsao_entrega", "data_estimada",
      "entrega_prevista", "data_limite", "dt_prevista", "dt_prometida", "previsao",
      "order_estimated_delivery_date", "data_previsao_entrega"]),
    ("data_entrega", "Data em que foi entregue", True,
     ["data_entrega", "entregue_em", "data_entregue", "dt_entrega", "data_recebimento", "recebido_em",
      "entrega_realizada", "data_entrega_realizada", "order_delivered_customer_date", "entregue",
      "chegou", "chegada", "data_chegada", "dt_chegada"]),
    ("valor_produtos", "Valor dos produtos (sem frete)", True,
     ["valor", "valor_produtos", "valor_produto", "valor_total", "total_produtos", "price", "preco",
      "valor_pedido", "subtotal", "valor_itens", "valor_mercadoria"]),
    ("valor_frete", "Valor do frete", True,
     ["frete", "valor_frete", "frete_total", "freight_value", "custo_frete", "vl_frete"]),
    ("n_itens", "Quantidade de itens", True,
     ["itens", "n_itens", "qtd_itens", "quantidade", "quantidade_itens", "qtd", "qtde", "qtde_itens",
      "num_itens", "no_itens"]),
    ("n_vendedores", "Quantidade de vendedores (ou lojas) no pedido", False,
     ["vendedores", "n_vendedores", "qtd_vendedores", "quantidade_vendedores", "qtde_vendedores", "lojas",
      "qtd_lojas", "sellers"]),
    ("n_produtos", "Quantidade de produtos diferentes", False,
     ["produtos", "n_produtos", "produtos_diferentes", "qtd_produtos", "qtde_produtos"]),
    ("estado_cliente", "Estado do cliente (UF)", False,
     ["estado_cliente", "uf_cliente", "uf", "estado", "customer_state", "uf_destino", "estado_destino",
      "uf_entrega", "estado_entrega"]),
    ("estado_vendedor", "Estado do vendedor (UF)", False,
     ["estado_vendedor", "uf_vendedor", "seller_state", "uf_origem", "estado_origem", "uf_loja",
      "estado_loja"]),
    ("forma_pagamento", "Forma de pagamento", False,
     ["forma_pagamento", "pagamento", "payment_type", "meio_pagamento", "tipo_pagamento", "forma_pgto",
      "pgto"]),
    ("n_parcelas", "Número de parcelas", False,
     ["parcelas", "n_parcelas", "qtd_parcelas", "num_parcelas", "payment_installments", "no_parcelas"]),
    ("n_formas_pagamento", "Quantidade de formas de pagamento usadas", False,
     ["formas_pagamento", "n_formas_pagamento", "qtd_formas_pagamento"]),
    ("categoria", "Categoria do produto", False,
     ["categoria", "categoria_produto", "product_category_name", "departamento", "secao", "categ"]),
    ("data_aprovacao", "Data/hora da aprovação do pagamento", False,
     ["data_aprovacao", "aprovado_em", "order_approved_at", "data_pagamento", "pago_em", "dt_aprovacao"]),
    ("data_envio", "Data de envio (entrega à transportadora)", False,
     ["data_envio", "enviado_em", "data_postagem", "data_despacho", "dt_envio", "postado_em",
      "order_delivered_carrier_date"]),
    ("peso_g", "Peso total em gramas", False,
     ["peso", "peso_g", "peso_total", "peso_total_g", "peso_gramas"]),
    ("fotos", "Média de fotos nos anúncios", False,
     ["fotos", "n_fotos", "fotos_media", "qtd_fotos"]),
    ("tam_descricao", "Tamanho da descrição dos anúncios (caracteres)", False,
     ["tam_descricao", "tamanho_descricao", "tam_descricao_media"]),
]
DESCRICAO = {c[0]: c[1] for c in CAMPOS}
OBRIGATORIOS = [c[0] for c in CAMPOS if c[2]]

UFS = {"acre": "AC", "alagoas": "AL", "amapa": "AP", "amazonas": "AM", "bahia": "BA", "ceara": "CE",
       "distrito federal": "DF", "espirito santo": "ES", "goias": "GO", "maranhao": "MA",
       "mato grosso": "MT", "mato grosso do sul": "MS", "minas gerais": "MG", "para": "PA",
       "paraiba": "PB", "parana": "PR", "pernambuco": "PE", "piaui": "PI", "rio de janeiro": "RJ",
       "rio grande do norte": "RN", "rio grande do sul": "RS", "rondonia": "RO", "roraima": "RR",
       "santa catarina": "SC", "sao paulo": "SP", "sergipe": "SE", "tocantins": "TO"}


# ---------------------------------------------------------------------------
# Utilidades
# ---------------------------------------------------------------------------
def caminho_recurso(relativo):
    """Acha arquivos embutidos no .exe (PyInstaller) ou ao lado do script."""
    base = getattr(sys, "_MEIPASS", os.path.dirname(os.path.abspath(__file__)))
    return os.path.join(base, relativo)


def sem_acento(texto):
    return unicodedata.normalize("NFKD", str(texto)).encode("ascii", "ignore").decode("ascii")


def normaliza_nome(nome):
    """'Nº do Pedido' -> 'no_pedido'; 'Valor (R$)' -> 'valor'."""
    t = sem_acento(nome).lower().replace("r$", " ")
    partes = [p for p in re.split(r"[^a-z0-9]+", t) if p and p not in {"do", "da", "de", "dos", "das", "em"}]
    return "_".join(partes)


SINONIMOS = {}
for _chave, _desc, _obrig, _nomes in CAMPOS:
    for _n in _nomes + [_chave, _desc]:
        SINONIMOS.setdefault(normaliza_nome(_n), _chave)


def arquivo_config():
    pasta = os.path.join(os.environ.get("APPDATA") or os.path.expanduser("~"), "AvaliacaoRisco")
    return os.path.join(pasta, "colunas.json")


def carrega_config():
    try:
        with open(arquivo_config(), encoding="utf-8") as f:
            return json.load(f)
    except (OSError, ValueError):
        return {}


def salva_config(config):
    try:
        os.makedirs(os.path.dirname(arquivo_config()), exist_ok=True)
        with open(arquivo_config(), "w", encoding="utf-8") as f:
            json.dump(config, f, ensure_ascii=False, indent=2)
    except OSError:
        pass  # sem permissão: o programa funciona, só não lembra da escolha


def assinatura(colunas):
    return "|".join(sorted(normaliza_nome(c) for c in colunas))


# ---------------------------------------------------------------------------
# Leitura e conversão de valores
# ---------------------------------------------------------------------------
def ler_planilha(caminho):
    ext = os.path.splitext(caminho)[1].lower()
    if ext in (".xlsx", ".xlsm"):
        dados = pd.read_excel(caminho, sheet_name=0)
    elif ext == ".xls":
        raise ValueError("Arquivos .xls (Excel antigo) não são aceitos.\n"
                         "Abra no Excel e salve como .xlsx ou .csv.")
    else:
        dados = pd.read_csv(caminho, sep=None, engine="python", encoding="utf-8-sig", dtype=str)
    dados = dados.dropna(how="all")
    dados.columns = [str(c).strip() for c in dados.columns]
    if dados.empty:
        raise ValueError("A planilha está vazia.")
    return dados


def para_numero(serie):
    """Aceita 1234.56, 1234,56, 1.234,56, R$ 1.234,56 e células vazias."""
    if pd.api.types.is_numeric_dtype(serie):
        return serie.astype(float)

    def conv(v):
        if pd.isna(v):
            return np.nan
        t = re.sub(r"[^0-9,.\-]", "", str(v))
        if not t:
            return np.nan
        if "," in t and "." in t:
            t = t.replace(".", "").replace(",", ".") if t.rfind(",") > t.rfind(".") else t.replace(",", "")
        elif "," in t:
            t = t.replace(",", ".")
        try:
            return float(t)
        except ValueError:
            return np.nan
    return serie.map(conv).astype(float)


def para_data(serie):
    """Aceita datas do Excel, 2018-08-02, 02/08/2018 e 02/08/2018 14:30."""
    if pd.api.types.is_datetime64_any_dtype(serie):
        return serie
    texto = serie.astype("string").str.strip()
    iso = texto.str.match(r"^\d{4}-\d{2}-\d{2}", na=False)
    resultado = pd.Series(pd.NaT, index=serie.index, dtype="datetime64[ns]")
    if iso.any():
        resultado[iso] = pd.to_datetime(texto[iso], errors="coerce", format="mixed")
    resto = ~iso & texto.notna()
    if resto.any():
        resultado[resto] = pd.to_datetime(texto[resto], errors="coerce", dayfirst=True, format="mixed")
    return resultado


def para_uf(serie):
    def conv(v):
        if pd.isna(v):
            return np.nan
        t = sem_acento(v).strip().lower()
        if len(t) == 2:
            return t.upper()
        return UFS.get(t, np.nan)
    return serie.map(conv)


def para_pagamento(serie):
    def conv(v):
        if pd.isna(v):
            return np.nan
        t = sem_acento(v).lower()
        if "debit" in t:
            return "debit_card"
        if "credit" in t or "cartao" in t or "credito" in t:
            return "credit_card"
        if "boleto" in t:
            return "boleto"
        if "voucher" in t or "vale" in t or "cupom" in t:
            return "voucher"
        return normaliza_nome(t) or np.nan  # ex.: pix (desconhecido para o modelo)
    return serie.map(conv)


def para_categoria(serie, conhecidas):
    """Aproxima a categoria da empresa das categorias que o modelo conhece."""
    chave = {normaliza_nome(c).replace("_", ""): c for c in conhecidas}
    cache = {}

    def conv(v):
        if pd.isna(v):
            return np.nan
        n = normaliza_nome(v)
        if n in cache:
            return cache[n]
        if n in conhecidas:
            cache[n] = n
        else:
            compacto = n.replace("_", "")
            parecidas = difflib.get_close_matches(compacto, list(chave), n=1, cutoff=0.8)
            cache[n] = chave[parecidas[0]] if parecidas else n
        return cache[n]
    return serie.map(conv)


# ---------------------------------------------------------------------------
# Mapeamento de colunas
# ---------------------------------------------------------------------------
def sugere_mapeamento(colunas, lembrado=None):
    """Retorna {campo: coluna_da_planilha} com base nos nomes comuns e em escolhas anteriores."""
    mapa = {}
    usadas = set()
    if lembrado:
        for col, campo in lembrado.items():
            if col in colunas and campo not in mapa:
                mapa[campo] = col
                usadas.add(col)
    for col in colunas:
        campo = SINONIMOS.get(normaliza_nome(col))
        if campo and campo not in mapa and col not in usadas:
            mapa[campo] = col
            usadas.add(col)
    return mapa


def obter_mapeamento(dados, perguntar):
    """Decide qual coluna é cada informação. `perguntar` abre a janela de escolha (ou None nos testes)."""
    config = carrega_config()
    sig = assinatura(dados.columns)
    lembrado = config.get(sig)
    mapa = sugere_mapeamento(list(dados.columns), lembrado)

    faltando = [c for c in OBRIGATORIOS if c not in mapa]
    if lembrado is None or faltando:
        if perguntar is None:
            if faltando:
                raise ValueError("Não encontrei as colunas de: " + ", ".join(DESCRICAO[c] for c in faltando))
            return mapa
        mapa = perguntar(dados, mapa)
        if mapa is None:
            return None  # usuário cancelou
        config[sig] = {col: campo for campo, col in mapa.items()}
        salva_config(config)
    return mapa


# ---------------------------------------------------------------------------
# Cálculo das 24 colunas do modelo a partir da planilha simples
# ---------------------------------------------------------------------------
def calcula_colunas(dados, mapa, categorias_conhecidas):
    def col(campo):
        return dados[mapa[campo]] if campo in mapa else pd.Series(np.nan, index=dados.index)

    compra = para_data(col("data_compra"))
    prometida = para_data(col("data_prometida"))
    entrega = para_data(col("data_entrega"))
    aprovacao = para_data(col("data_aprovacao"))
    envio = para_data(col("data_envio"))

    f = pd.DataFrame(index=dados.index)
    f["dias_entrega"] = (entrega - compra).dt.days
    f["dias_prometidos"] = (prometida - compra).dt.days
    f["dias_atraso"] = (entrega - prometida).dt.days
    f["atrasado"] = np.where(f["dias_atraso"].isna(), np.nan, (f["dias_atraso"] > 0).astype(float))
    f["horas_aprovacao"] = (aprovacao - compra).dt.total_seconds() / 3600
    f["dias_ate_transportadora"] = (envio - compra).dt.days
    f["mes_compra"] = compra.dt.month
    f["dia_semana_compra"] = compra.dt.dayofweek

    f["n_itens"] = para_numero(col("n_itens")).round()
    f["n_produtos"] = para_numero(col("n_produtos")).round().fillna(f["n_itens"])
    f["n_vendedores"] = para_numero(col("n_vendedores")).round().fillna(1)
    f["valor_total"] = para_numero(col("valor_produtos"))
    f["frete_total"] = para_numero(col("valor_frete"))
    f["frete_sobre_valor"] = f["frete_total"] / f["valor_total"].replace(0, np.nan)
    f["peso_total_g"] = para_numero(col("peso_g"))
    f["fotos_media"] = para_numero(col("fotos"))
    f["tam_descricao_media"] = para_numero(col("tam_descricao"))
    f["n_parcelas"] = para_numero(col("n_parcelas")).fillna(1)
    f["n_formas_pagamento"] = para_numero(col("n_formas_pagamento")).fillna(1)

    f["categoria"] = para_categoria(col("categoria"), categorias_conhecidas)
    f["estado_vendedor"] = para_uf(col("estado_vendedor"))
    f["estado_cliente"] = para_uf(col("estado_cliente"))
    f["forma_pagamento"] = para_pagamento(col("forma_pagamento"))
    ambos = f["estado_cliente"].notna() & f["estado_vendedor"].notna()
    f["mesmo_estado"] = np.where(ambos, (f["estado_cliente"] == f["estado_vendedor"]).astype(float), np.nan)
    for c in CATEGORICAS:  # texto, com np.nan nos vazios (igual ao treino)
        f[c] = f[c].astype(object).where(f[c].notna(), np.nan)

    # Pedidos que não dá para analisar
    motivo = pd.Series("", index=dados.index)
    motivo[prometida.isna()] = "Data prometida de entrega vazia ou em formato não reconhecido"
    motivo[compra.isna()] = "Data da compra vazia ou em formato não reconhecido"
    motivo[entrega.isna()] = "Sem data de entrega (pedido ainda não entregue?) ou formato não reconhecido"
    motivo[(motivo == "") & (f["dias_entrega"] < 0)] = "Data de entrega anterior à data da compra"
    motivo[(motivo == "") & f["n_itens"].isna()] = "Quantidade de itens vazia"
    return f, motivo


# ---------------------------------------------------------------------------
# Explicação do risco
# ---------------------------------------------------------------------------
def motivo_e_acao(p):
    motivos, acoes = [], []
    atraso = p["dias_atraso"]
    if pd.notna(atraso) and atraso > 0:
        motivos.append(f"Entregue com {int(round(atraso))} dia(s) de atraso")
        acoes.append("pedir desculpas pelo atraso e oferecer um cupom")
    if pd.notna(p["n_vendedores"]) and p["n_vendedores"] >= 2:
        motivos.append(f"{int(p['n_vendedores'])} vendedores: pode ter chegado só parte do pedido")
        acoes.append("confirmar se o cliente recebeu todos os itens")
    elif pd.notna(p["n_itens"]) and p["n_itens"] >= 3:
        motivos.append(f"Pedido com {int(p['n_itens'])} itens")
        acoes.append("confirmar se todos os itens chegaram em ordem")
    if not motivos:
        motivos.append("Combinação de fatores (prazo, frete, produto)")
        acoes.append("verificar se está satisfeito com a compra")
    return " | ".join(motivos), "Contatar o cliente: " + "; ".join(acoes)


# ---------------------------------------------------------------------------
# Processamento principal
# ---------------------------------------------------------------------------
def processar(entrada, perguntar=None):
    """Lê a planilha, calcula o risco e grava o Excel. Retorna um resumo (dict) ou None se cancelado."""
    modelo = joblib.load(caminho_recurso(os.path.join("models", "modelo_avaliacao_negativa.joblib")))
    categorias = [c for c in modelo.named_steps["preparacao"].named_transformers_["cat"].categories_[0]
                  if isinstance(c, str)]
    dados = ler_planilha(entrada)

    formato_completo = all(c in dados.columns for c in COLUNAS_MODELO)
    if formato_completo:
        id_col = "order_id" if "order_id" in dados.columns else None
        feats = dados[COLUNAS_MODELO].copy()
        for c in NUMERICAS:
            feats[c] = para_numero(feats[c])
        for c in CATEGORICAS:
            feats[c] = feats[c].astype(object).where(feats[c].notna(), np.nan)
        nao_analisar = pd.Series("", index=dados.index)
        nao_analisar[feats["dias_atraso"].isna() & feats["dias_entrega"].isna()] = "Sem informações de prazo"
        usadas = set(COLUNAS_MODELO)
        faltaram = []
    else:
        mapa = obter_mapeamento(dados, perguntar)
        if mapa is None:
            return None
        id_col = mapa.get("pedido")
        feats, nao_analisar = calcula_colunas(dados, mapa, categorias)
        usadas = set()  # no formato simples, todas as colunas originais vão para o resultado
        faltaram = [DESCRICAO[c] for c, _d, obrig, _n in CAMPOS if c not in mapa and not obrig]

    ok = nao_analisar == ""
    if not ok.any():
        raise ValueError("Nenhum pedido pôde ser analisado.\n\nPrimeiros problemas encontrados:\n- "
                         + "\n- ".join(nao_analisar.head(3)))

    prob = pd.Series(np.nan, index=dados.index)
    prob[ok] = modelo.predict_proba(feats.loc[ok, modelo.feature_names_in_])[:, 1]

    explic = feats.loc[ok].apply(motivo_e_acao, axis=1, result_type="expand")
    identificacao = dados[id_col] if id_col else pd.Series(range(2, len(dados) + 2), index=dados.index)

    resultado = pd.DataFrame({
        "pedido": identificacao,
        "risco": prob.round(3),
        "nivel": ["ALTO" if p >= CORTE else "MÉDIO" if p >= CORTE_MEDIO else "BAIXO" for p in prob],
        "motivo": explic[0].reindex(dados.index),
        "acao_sugerida": explic[1].reindex(dados.index),
        "dias_de_atraso": feats["dias_atraso"],
        "itens": feats["n_itens"],
        "vendedores": feats["n_vendedores"],
    }, index=dados.index)
    baixo = resultado["nivel"] == "BAIXO"
    resultado.loc[baixo, "motivo"] = "-"
    resultado.loc[baixo, "acao_sugerida"] = "Nenhuma ação necessária"

    # Colunas originais da planilha (nome, telefone, datas...) vão junto
    extras = [c for c in dados.columns if c != id_col and c not in usadas and c not in resultado.columns]
    resultado = pd.concat([resultado, dados[extras]], axis=1)

    analisados = resultado[ok].sort_values("risco", ascending=False)
    atencao = analisados[analisados["nivel"] == "ALTO"]
    problemas = pd.concat([pd.DataFrame({"pedido": identificacao[~ok], "linha_na_planilha": (dados.index[~ok] + 2),
                                         "motivo": nao_analisar[~ok]}), dados.loc[~ok, extras]], axis=1)

    saida = os.path.splitext(entrada)[0] + "_com_risco.xlsx"
    with pd.ExcelWriter(saida, engine="openpyxl") as excel:
        if len(atencao):
            atencao.to_excel(excel, sheet_name="Precisam de atenção", index=False)
        else:
            pd.DataFrame({"mensagem": ["Nenhum pedido com risco alto nesta planilha."]}).to_excel(
                excel, sheet_name="Precisam de atenção", index=False)
        analisados.to_excel(excel, sheet_name="Todos os pedidos", index=False)
        if len(problemas):
            problemas.to_excel(excel, sheet_name="Não analisados", index=False)
        for aba in excel.book.worksheets:
            formatar(aba)

    return {"saida": saida, "analisados": len(analisados), "atencao": len(atencao),
            "nao_analisados": len(problemas), "faltaram": faltaram}


def formatar(aba):
    """Cabeçalho destacado, larguras ajustadas, datas legíveis e cores por nível de risco."""
    cores = {n: PatternFill("solid", fgColor=c) for n, c in
             {"ALTO": "F8CBAD", "MÉDIO": "FFE699", "BAIXO": "C6EFCE"}.items()}
    cabecalho = [c.value for c in aba[1]]
    fundo, fonte = PatternFill("solid", fgColor="305496"), Font(bold=True, color="FFFFFF")
    alinhamento = Alignment(vertical="center", wrap_text=True)
    for c in aba[1]:
        c.font, c.fill, c.alignment = fonte, fundo, alinhamento

    larguras = [len(str(n or "")) for n in cabecalho]
    col_nivel = cabecalho.index("nivel") if "nivel" in cabecalho else None
    col_risco = cabecalho.index("risco") if "risco" in cabecalho else None
    for i, linha in enumerate(aba.iter_rows(min_row=2), start=2):
        cor = cores.get(linha[col_nivel].value) if col_nivel is not None else None
        for j, cel in enumerate(linha):
            if cor:
                cel.fill = cor
            if hasattr(cel.value, "year"):
                cel.number_format = "dd/mm/yyyy"
                larguras[j] = max(larguras[j], 10)
            elif i < 500:  # largura estimada pelas primeiras linhas
                larguras[j] = max(larguras[j], len(str(cel.value if cel.value is not None else "")))
        if col_risco is not None:
            linha[col_risco].number_format = "0%"
    for j, largura in enumerate(larguras):
        aba.column_dimensions[aba.cell(row=1, column=j + 1).column_letter].width = min(largura + 2, 60)
    aba.freeze_panes = "B2"


# ---------------------------------------------------------------------------
# Interface (janelas)
# ---------------------------------------------------------------------------
def janela_colunas(raiz, dados, sugestao):
    """Pergunta qual coluna da planilha corresponde a cada informação."""
    import tkinter as tk
    from tkinter import messagebox, ttk

    NADA = "(não tenho esta informação)"
    opcoes = [NADA] + list(dados.columns)
    resposta = {"mapa": None}

    jan = tk.Toplevel(raiz)
    jan.title("Avaliação de risco: colunas da planilha")
    jan.resizable(False, False)
    tk.Label(jan, justify="left", wraplength=640, padx=12, pady=10,
             text="Indique qual coluna da sua planilha corresponde a cada informação. "
                  "As marcadas com * são obrigatórias. As outras melhoram a previsão, mas podem ficar "
                  "em branco.\nO programa vai lembrar desta escolha nas próximas vezes.").grid(
        row=0, column=0, columnspan=3, sticky="w")

    combos = {}
    for i, (chave, desc, obrig, _) in enumerate(CAMPOS, start=1):
        tk.Label(jan, text=desc + (" *" if obrig else ""), anchor="w",
                 font=("Segoe UI", 9, "bold" if obrig else "normal")).grid(row=i, column=0, sticky="w", padx=12, pady=2)
        cb = ttk.Combobox(jan, values=opcoes, state="readonly", width=32)
        cb.set(sugestao.get(chave, NADA))
        cb.grid(row=i, column=1, padx=6, pady=2)
        exemplo = tk.Label(jan, fg="gray40", anchor="w", width=28)
        exemplo.grid(row=i, column=2, sticky="w", padx=(0, 12))

        def mostra_exemplo(_evento=None, cb=cb, lbl=exemplo):
            col = cb.get()
            if col == NADA:
                lbl.config(text="")
            else:
                vals = dados[col].dropna()
                if not len(vals):
                    lbl.config(text="(coluna vazia)")
                    return
                v = vals.iloc[0]
                v = v.strftime("%d/%m/%Y") if hasattr(v, "strftime") else str(v)
                lbl.config(text="ex.: " + v[:25])
        cb.bind("<<ComboboxSelected>>", mostra_exemplo)
        mostra_exemplo()
        combos[chave] = cb

    def confirmar():
        mapa = {k: cb.get() for k, cb in combos.items() if cb.get() != NADA}
        faltando = [DESCRICAO[k] for k in OBRIGATORIOS if k not in mapa]
        if faltando:
            messagebox.showwarning("Avaliação de risco", "Escolha a coluna de:\n\n- " + "\n- ".join(faltando),
                                   parent=jan)
            return
        repetidas = {c for c in mapa.values() if list(mapa.values()).count(c) > 1}
        if repetidas:
            messagebox.showwarning("Avaliação de risco",
                                   "Cada coluna só pode ser usada uma vez:\n\n- " + "\n- ".join(repetidas),
                                   parent=jan)
            return
        resposta["mapa"] = mapa
        jan.destroy()

    botoes = tk.Frame(jan, pady=10)
    botoes.grid(row=len(CAMPOS) + 1, column=0, columnspan=3)
    ttk.Button(botoes, text="Continuar", command=confirmar).pack(side="left", padx=6)
    ttk.Button(botoes, text="Cancelar", command=jan.destroy).pack(side="left", padx=6)

    jan.protocol("WM_DELETE_WINDOW", jan.destroy)
    jan.attributes("-topmost", True)
    try:
        jan.wait_visibility()
        jan.grab_set()
    except tk.TclError:
        pass
    raiz.wait_window(jan)
    return resposta["mapa"]


def main():
    import tkinter as tk
    from tkinter import filedialog, messagebox

    raiz = tk.Tk()
    raiz.withdraw()

    if len(sys.argv) > 1:          # planilha arrastada para cima do programa
        entrada = sys.argv[1]
    else:                          # dois cliques: pede a planilha
        entrada = filedialog.askopenfilename(
            title="Escolha a planilha com os pedidos entregues",
            filetypes=[("Planilhas", "*.csv *.xlsx *.xlsm"), ("CSV", "*.csv"), ("Excel", "*.xlsx *.xlsm"),
                       ("Todos os arquivos", "*.*")])
    if not entrada:
        return

    try:
        r = processar(entrada, perguntar=lambda dados, sug: janela_colunas(raiz, dados, sug))
    except PermissionError:
        messagebox.showerror("Avaliação de risco", "Feche o arquivo de resultado no Excel e tente de novo.")
        return
    except Exception as erro:
        messagebox.showerror("Avaliação de risco", f"Não foi possível processar a planilha.\n\n{erro}")
        return
    if r is None:
        return

    texto = f"{r['analisados']} pedidos analisados\n{r['atencao']} precisam de atenção (risco alto)\n"
    if r["nao_analisados"]:
        texto += f"{r['nao_analisados']} não puderam ser analisados (veja a aba \"Não analisados\")\n"
    if r["faltaram"]:
        texto += ("\nInformações opcionais que não estavam na planilha (a previsão funciona sem elas, "
                  "mas fica um pouco menos precisa):\n- " + "\n- ".join(r["faltaram"]) + "\n")
    texto += f"\nResultado salvo em:\n{r['saida']}\n\nAbrir o resultado agora?"
    if messagebox.askyesno(f"Avaliação de risco v{VERSAO}", texto) and hasattr(os, "startfile"):
        os.startfile(r["saida"])


if __name__ == "__main__":
    main()
